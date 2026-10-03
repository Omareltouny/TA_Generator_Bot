import io
import zipfile

import docx
import pytest

from bot.services import renderer as R

META = R.DocMeta("CS 2920", "Data Structures", "lab", 3, 4, "Binary Trees")
MD = """### Task 1
Compute $x^2 + y^2 = z^2$ and
$$\\int_0^1 x\\,dx = \\frac{1}{2}$$

```java
public class A {
    int x;
}
```
| a | b |
|---|---|
| 1 | 2 |
"""


def tiny_png() -> bytes:
    import base64
    return base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


async def test_docx_native_equations_monospace_header():
    data = await R.render_docx(MD, META, answer_key=True, logo=tiny_png())
    z = zipfile.ZipFile(io.BytesIO(data))
    xml = z.read("word/document.xml").decode()
    assert "<m:oMath" in xml and "<pic:pic" not in xml  # native OMML, not images
    d = docx.Document(io.BytesIO(data))
    code = [p for p in d.paragraphs if p.style.name == "Source Code"]
    assert code and code[0].style.font.name == "Consolas" and "    int x;" in "\n".join(p.text for p in code)
    assert any("Answer Key" in p.text for p in d.paragraphs)
    hdr = d.sections[0].header
    assert "CS 2920" in hdr.paragraphs[0].text and "Week 4" in hdr.paragraphs[0].text
    assert "word/media/" in " ".join(z.namelist())  # logo embedded
    assert "PAGE" in z.read("word/footer1.xml").decode() or any("PAGE" in z.read(n).decode() for n in z.namelist() if "footer" in n)


async def test_student_title_has_no_answer_key_and_tex():
    d = docx.Document(io.BytesIO(await R.render_docx(MD, META, answer_key=False)))
    assert not any("Answer Key" in p.text for p in d.paragraphs)
    tex = (await R.render_tex(MD, META, answer_key=False)).decode()
    assert "\\frac{1}{2}" in tex and "\\begin{document}" in tex


async def test_bad_logo_does_not_break():
    assert await R.render_docx(MD, META, answer_key=False, logo=b"garbage")


def test_template_validation():
    assert R.validate_template(R.DEFAULT_REFERENCE.read_bytes()) == []
    with pytest.raises(ValueError):
        R.validate_template(b"nope")
    d = docx.Document()
    b = io.BytesIO(); d.save(b)
    assert "Source Code" in R.validate_template(b.getvalue())


def test_filenames():
    assert META.filename("student", "docx") == "CS_2920_lab3_student.docx"
