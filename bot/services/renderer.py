"""Markdown (+LaTeX math) -> .docx via pandoc (native OMML equations), python-docx post-processing, .tex export."""
from __future__ import annotations

import asyncio
import io
import os
import pathlib
import tempfile
from dataclasses import dataclass

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

DEFAULT_REFERENCE = pathlib.Path(__file__).resolve().parent.parent / "templates" / "reference.docx"
PANDOC_FROM = "markdown+tex_math_dollars+pipe_tables+fenced_code_blocks+backtick_code_blocks-smart"
MONO = "Consolas"
REQUIRED_STYLES = ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Body Text", "Source Code"]


class RenderError(Exception):
    pass


@dataclass
class DocMeta:
    code: str
    course_name: str
    type: str
    seq: int
    week: int | None
    title: str

    @property
    def doc_title(self) -> str:
        return f"{self.type.title()} {self.seq}: {self.title}"

    @property
    def header_text(self) -> str:
        left = " - ".join(x for x in [self.code, self.course_name] if x)
        right = f"{self.type.title()} {self.seq}" + (f" | Week {self.week}" if self.week else "")
        return f"{left}\t{right}" if left else right

    def filename(self, kind: str, ext: str) -> str:
        base = "".join(c if c.isalnum() else "_" for c in (self.code or "course")).strip("_")
        return f"{base}_{self.type}{self.seq}_{kind}.{ext}"


async def _pandoc(args: list[str], md: str) -> bytes:
    try:
        proc = await asyncio.create_subprocess_exec("pandoc", *args, stdin=asyncio.subprocess.PIPE,
                                                    stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    except FileNotFoundError:
        raise RenderError("pandoc is not installed")
    out, err = await asyncio.wait_for(proc.communicate(md.encode()), timeout=90)
    if proc.returncode != 0:
        raise RenderError(f"pandoc failed: {err.decode()[:300]}")
    return out


def validate_template(data: bytes) -> list[str]:
    """Return missing required styles (empty = fine). Raises ValueError if not a docx."""
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise ValueError("not a valid .docx file")
    have = {s.name for s in d.styles}
    return [s for s in REQUIRED_STYLES if s not in have]


def _set_mono(style) -> None:
    style.font.name = MONO
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), MONO)


def _field(par, instr: str) -> None:
    run = par.add_run()
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def postprocess(docx_bytes: bytes, meta: DocMeta, logo: bytes | None) -> bytes:
    d = docx.Document(io.BytesIO(docx_bytes))
    names = {s.name: s for s in d.styles}
    for n in ("Source Code", "Verbatim Char"):  # guarantee monospace even for custom templates
        if n in names:
            _set_mono(names[n])
    sec = d.sections[0]
    hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    # pandoc output may omit page geometry; fill in US Letter / 1in margins so layout is deterministic
    sec.page_width = sec.page_width or Inches(8.5)
    sec.page_height = sec.page_height or Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        if getattr(sec, side) is None:
            setattr(sec, side, Inches(1))
    width = sec.page_width - sec.left_margin - sec.right_margin
    hp.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)  # right-aligned tab
    if logo:
        try:
            hp.add_run().add_picture(io.BytesIO(logo), height=Cm(1.4))
            hp.add_run("   ")
        except Exception:
            pass  # an unreadable logo must never break document generation
    r = hp.add_run(meta.header_text)
    r.font.size = Pt(9)
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = 1
    fp.add_run("Page ")
    _field(fp, "PAGE")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def render_docx(md: str, meta: DocMeta, *, answer_key: bool, logo: bytes | None = None,
                      template: bytes | None = None) -> bytes:
    title = meta.doc_title + (" - Answer Key" if answer_key else "")
    fd, ref = tempfile.mkstemp(suffix=".docx")  # scratch only; the DB is the source of truth
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(template or DEFAULT_REFERENCE.read_bytes())
        args = ["-f", PANDOC_FROM, "-t", "docx", "--no-highlight", f"--reference-doc={ref}", "-M", f"title={title}"]
        if meta.course_name:
            args += ["-M", f"subtitle={meta.course_name}"]
        out = await _pandoc(args, md)
    finally:
        os.unlink(ref)
    return postprocess(out, meta, logo)


async def render_pair(student_md: str, key_md: str, meta: DocMeta, logo=None, template=None) -> tuple[bytes, bytes]:
    return (await render_docx(student_md, meta, answer_key=False, logo=logo, template=template),
            await render_docx(key_md, meta, answer_key=True, logo=logo, template=template))


async def render_tex(md: str, meta: DocMeta, *, answer_key: bool) -> bytes:
    title = meta.doc_title + (" - Answer Key" if answer_key else "")
    return await _pandoc(["-f", PANDOC_FROM, "-t", "latex", "-s", "-M", f"title={title}"], md)
