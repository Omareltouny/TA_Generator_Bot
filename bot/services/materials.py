"""Text extraction for uploaded material, safe zip handling, chunking and retrieval."""
from __future__ import annotations

import io
import math
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from xml.etree import ElementTree

from bot.services.pdf_text import extract_pdf

CODE_EXT = {".py", ".java", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".kt", ".js", ".ts", ".go", ".rs", ".sh",
            ".sql", ".m", ".r", ".swift", ".rb", ".php", ".asm", ".s", ".v", ".vhd", ".tex", ".csv", ".json", ".xml"}
TEXT_EXT = {".txt", ".md"} | CODE_EXT
ACCEPTED = {".pdf", ".docx", ".pptx"} | TEXT_EXT
MAX_TEXT_PER_FILE = 300_000
ZIP_MAX_FILES, ZIP_MAX_TOTAL, ZIP_MAX_RATIO = 300, 150 * 1024 * 1024, 200


def ext_of(name: str) -> str:
    return "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""


@dataclass
class Extracted:
    text: str = ""
    warning: str | None = None


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[0]))
        for n in slides:
            root = ElementTree.fromstring(z.read(n))
            texts = [e.text for e in root.iter() if e.tag.endswith("}t") and e.text]
            out.append(f"--- slide {re.findall(r'[0-9]+', n)[0]} ---\n" + " ".join(texts))
    return "\n".join(out)


def extract_text(name: str, data: bytes) -> Extracted:
    ext = ext_of(name)
    if ext not in ACCEPTED:
        raise ValueError(f"unsupported type {ext or '(none)'}")
    warn = None
    if ext == ".pdf":
        e = extract_pdf(data)
        text = e.text
        if e.empty_pages and len(e.empty_pages) == e.pages:
            raise ValueError("PDF has no text layer (scanned images); OCR is not supported")
        if e.empty_pages:
            warn = f"{len(e.empty_pages)} page(s) had no text layer and were skipped"
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext == ".pptx":
        text = _pptx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")
    if len(text) > MAX_TEXT_PER_FILE:
        text, warn = text[:MAX_TEXT_PER_FILE], f"truncated to {MAX_TEXT_PER_FILE:,} characters"
    return Extracted(text, warn)


@dataclass
class ZipResult:
    files: list[tuple[str, bytes]] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)


def read_zip(data: bytes) -> ZipResult:
    """In-memory only (nothing is written to disk, so path traversal cannot escape), with bomb guards."""
    res = ZipResult()
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ValueError("not a valid zip file")
    infos = [i for i in z.infolist() if not i.is_dir()]
    if len(infos) > ZIP_MAX_FILES:
        raise ValueError(f"zip has too many files ({len(infos)} > {ZIP_MAX_FILES}); split it")
    if sum(i.file_size for i in infos) > ZIP_MAX_TOTAL:
        raise ValueError("zip expands to too much data (possible zip bomb)")
    for i in infos:
        n = i.filename
        base = n.rsplit("/", 1)[-1]
        if n.startswith(("/", "\\")) or ".." in n.replace("\\", "/").split("/") or base.startswith(".") or "__MACOSX" in n:
            res.skipped.append(f"{n} (unsafe/hidden path)")
        elif ext_of(base) not in ACCEPTED:
            res.skipped.append(f"{n} (unsupported type)")
        elif i.compress_size and i.file_size / max(i.compress_size, 1) > ZIP_MAX_RATIO:
            res.skipped.append(f"{n} (suspicious compression ratio)")
        else:
            res.files.append((n, z.read(i)))
    return res


# ---- chunk + retrieve ---------------------------------------------------
def chunk_text(text: str, size: int = 1200, overlap: int = 150) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks, i = [], 0
    while i < len(text):
        j = min(len(text), i + size)
        if j < len(text):  # prefer a paragraph/sentence boundary
            cut = max(text.rfind("\n", i + size // 2, j), text.rfind(". ", i + size // 2, j))
            j = cut + 1 if cut > 0 else j
        chunks.append(text[i:j].strip())
        if j >= len(text):
            break
        i = max(j - overlap, i + 1)
    return [c for c in chunks if c]


def _tok(s: str) -> list[str]:
    return [w for w in re.findall(r"[a-z0-9_+#]{3,}", s.lower())]


def retrieve(sources: list[tuple[str, str]], query: str, k: int = 4, max_chars: int = 4500) -> list[str]:
    """BM25-lite over chunks of all sources. Returns "[file] excerpt" strings, best first."""
    chunks = [(f, c) for f, t in sources for c in chunk_text(t)]
    if not chunks or not query.strip():
        return []
    toks = [Counter(_tok(c)) for _, c in chunks]
    df = Counter(w for t in toks for w in t)
    n, q = len(chunks), set(_tok(query))
    scored = []
    for idx, t in enumerate(toks):
        sc = sum((1 + math.log(t[w])) * math.log(1 + n / df[w]) for w in q if w in t)
        if sc > 0:
            scored.append((sc, idx))
    out, used = [], 0
    for _, idx in sorted(scored, reverse=True)[:k]:
        f, c = chunks[idx]
        if used + len(c) > max_chars:
            break
        out.append(f"[{f}] {c}")
        used += len(c)
    return out
